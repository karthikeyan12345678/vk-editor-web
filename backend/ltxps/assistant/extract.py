"""Reads text, tables and structural facts out of uploaded files (PDF, Word, PowerPoint, Excel, CSV, text)."""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".tsv", ".txt", ".md", ".json", ".html", ".htm", ".log"}
MAX_TABLE_ROWS = 500_000


class ExtractError(ValueError):
    pass


@dataclass
class Extracted:
    kind: str
    name: str
    text: str = ""
    units: list[str] = field(default_factory=list)  # paragraphs, slides, or pages in reading order
    tables: list[dict[str, Any]] = field(default_factory=list)  # {name, columns, rows}
    facts: dict[str, Any] = field(default_factory=dict)  # structural details used by the ATS checker
    warnings: list[str] = field(default_factory=list)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _sniff(sample: str) -> csv.Dialect | type[csv.excel]:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def read_csv_rows(path: Path, limit: int = MAX_TABLE_ROWS) -> tuple[list[str], list[list[str]], bool]:
    text = _decode(path.read_bytes())
    dialect = _sniff(text[:8192])
    reader = csv.reader(io.StringIO(text), dialect)
    rows: list[list[str]] = []
    truncated = False
    for r in reader:
        if not r or all(not c.strip() for c in r):
            continue
        rows.append(r)
        if len(rows) > limit:
            truncated = True
            rows.pop()
            break
    if not rows:
        raise ExtractError("The file has no rows.")
    cols = [c.strip() or f"column_{i + 1}" for i, c in enumerate(rows[0])]
    width = len(cols)
    body = [(r + [""] * width)[:width] for r in rows[1:]]
    return cols, body, truncated


def extract(path: str | Path) -> Extracted:
    p = Path(path)
    ext = p.suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractError(f"“{p.suffix or p.name}” files are not supported. Use: PDF, Word, PowerPoint, Excel, CSV or text files.")
    shown = re.sub(r"^\d{5}-", "", p.name)  # uploads are stored as 00012-name.ext; show the original name
    ex = Extracted(kind=ext.lstrip("."), name=shown)
    try:
        {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".csv": _csv, ".tsv": _csv}.get(ext, _plain)(p, ex)
    except ExtractError:
        raise
    except Exception as exc:  # noqa: BLE001 - corrupt files raise many library-specific errors
        raise ExtractError(f"“{p.name}” could not be read ({type(exc).__name__}). It may be damaged or password protected.") from exc
    if not ex.text and ex.tables:
        ex.text = "\n".join(" | ".join(t["columns"]) + "\n" + "\n".join(" | ".join(r) for r in t["rows"][:200]) for t in ex.tables)
    return ex


def _plain(p: Path, ex: Extracted) -> None:
    raw = _decode(p.read_bytes())
    if p.suffix.lower() in {".html", ".htm"}:
        raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
        raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    if p.suffix.lower() == ".json":
        try:
            raw = json.dumps(json.loads(raw), indent=2, ensure_ascii=False)
        except ValueError:
            ex.warnings.append("The JSON is not valid, so it is shown as plain text.")
    ex.text = raw
    ex.units = [u.strip() for u in re.split(r"\n\s*\n", raw) if u.strip()]


def _csv(p: Path, ex: Extracted) -> None:
    cols, rows, truncated = read_csv_rows(p)
    ex.tables = [{"name": Path(ex.name).stem, "columns": cols, "rows": rows}]
    if truncated:
        ex.warnings.append(f"Only the first {MAX_TABLE_ROWS:,} rows were read.")


def _xlsx(p: Path, ex: Extracted) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(str(p), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            if r is None or all(c is None or str(c).strip() == "" for c in r):
                continue
            rows.append(["" if c is None else (c.isoformat(sep=" ") if hasattr(c, "isoformat") else str(c)) for c in r])
            if len(rows) > MAX_TABLE_ROWS:
                ex.warnings.append(f"Sheet “{ws.title}”: only the first {MAX_TABLE_ROWS:,} rows were read.")
                break
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        cols = [c.strip() or f"column_{i + 1}" for i, c in enumerate(rows[0])]
        ex.tables.append({"name": ws.title, "columns": cols, "rows": rows[1:]})
    wb.close()
    if not ex.tables:
        raise ExtractError("The workbook has no data.")


def _pdf(p: Path, ex: Extracted) -> None:
    from pypdf import PdfReader

    r = PdfReader(str(p))
    if r.is_encrypted:
        try:
            if not r.decrypt(""):
                raise ExtractError("This PDF is password protected.")
        except Exception as exc:  # noqa: BLE001
            raise ExtractError("This PDF is password protected.") from exc
    pages = [re.sub(r"(?im)^\s*(page\s+)?\d+(\s*(of|/)\s*\d+)?\s*$", "", pg.extract_text() or "") for pg in r.pages]  # drop bare page numbers
    ex.units = pages
    ex.text = "\n\n".join(pages)
    images = 0
    for pg in r.pages:
        try:
            images += len(pg.images)
        except Exception:  # noqa: BLE001 - some PDFs have unreadable image dictionaries
            pass
    ex.facts = {"pages": len(pages), "images": images, "chars_per_page": (len(ex.text) / max(1, len(pages))), "text_extractable": len(ex.text.strip()) > 50}
    if not ex.facts["text_extractable"]:
        ex.warnings.append("No text could be read from this PDF. It may be a scan or an image; text recognition (OCR) is not included.")


def _docx(p: Path, ex: Extracted) -> None:
    from docx import Document
    from docx.oxml.ns import qn

    d = Document(str(p))
    paras = []
    heads = []
    for para in d.paragraphs:
        t = para.text.strip()
        if t:
            paras.append(t)
            if para.style is not None and para.style.name.lower().startswith(("heading", "title")):
                heads.append(t)
    for t in d.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        if rows:
            ex.tables.append({"name": f"table{len(ex.tables) + 1}", "columns": rows[0], "rows": rows[1:]})
    body = d.element.body
    textboxes = len(body.findall(".//" + qn("w:txbxContent")))
    cols = max([int(c.get(qn("w:num"), 1)) for c in body.iter(qn("w:cols"))] or [1])
    header_text = " ".join(par.text for s in d.sections for par in s.header.paragraphs).strip()
    footer_text = " ".join(par.text for s in d.sections for par in s.footer.paragraphs).strip()
    ex.units = paras
    table_text = [" ".join(c for r in [t["columns"]] + t["rows"] for c in r) for t in ex.tables]
    ex.text = "\n".join(paras)
    ex.facts = {
        "tables": len(d.tables), "images": len(d.inline_shapes), "textboxes": textboxes, "columns": cols, "headings": heads,
        "header_text": header_text, "footer_text": footer_text, "table_text": "\n".join(table_text),
        "fonts": sorted({r.font.name for par in d.paragraphs for r in par.runs if r.font.name}),
    }


def _pptx(p: Path, ex: Extracted) -> None:
    from pptx import Presentation

    prs = Presentation(str(p))
    slides = []
    for i, s in enumerate(prs.slides, start=1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = "\n".join(par.text for par in sh.text_frame.paragraphs if par.text.strip())
                if t:
                    parts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text for c in r.cells] for r in sh.table.rows]
                if rows:
                    ex.tables.append({"name": f"slide{i}-table", "columns": rows[0], "rows": rows[1:]})
                    parts.append("\n".join(" | ".join(r) for r in rows))
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        slides.append("\n".join(parts) + (f"\n[Notes] {notes}" if notes else ""))
    ex.units = slides
    ex.text = "\n\n".join(f"[Slide {i}]\n{t}" for i, t in enumerate(slides, start=1))
    ex.facts = {"slides": len(slides)}
