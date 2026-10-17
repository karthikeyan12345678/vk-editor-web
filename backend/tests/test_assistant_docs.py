import csv

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from ltxps.assistant.documents import build_file, parse_markdown
from ltxps.assistant.themes import catalog, contrast, document_palette, get_theme, recommend

MD = """# Quarterly report

Prepared for the **leadership** team.

## Highlights

- Revenue grew 18% year on year
- Churn fell to 2.1%
1. First numbered
2. Second numbered

## Numbers

| Region | Units | Revenue |
| --- | --- | --- |
| North | 120 | 4,500.50 |
| South | 95 | 3,100 |

## Next steps

Ship the beta, then review. Notes: not part of a doc.
"""


@pytest.fixture()
def doc():
    return parse_markdown(MD)


def test_pdf_has_text_and_pages(doc, tmp_path):
    r = build_file("pdf", doc, tmp_path, "report")
    text = "\n".join(p.extract_text() for p in PdfReader(r.path).pages)
    assert "Quarterly report" in text and "Revenue grew 18%" in text and "North" in text


def test_docx_has_headings_list_and_table(doc, tmp_path):
    r = build_file("docx", doc, tmp_path, "report")
    d = Document(r.path)
    assert any("Quarterly report" in p.text for p in d.paragraphs)
    assert len(d.tables) == 1 and d.tables[0].rows[1].cells[0].text == "North"
    assert any(p.text.startswith("Churn") for p in d.paragraphs)


def test_pptx_one_slide_per_heading_with_notes(tmp_path):
    md = "# Deck title\nSubtitle here\n\n## Problem\n- one\n- two\nNotes: say this\n\n## Solution\n- a\n- b\n- c\n- d\n- e\n- f\n- g\n"
    r = build_file("pptx", parse_markdown(md), tmp_path, "deck")
    p = Presentation(r.path)
    titles = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in p.slides]
    assert len(p.slides) >= 3
    assert any("Problem" in t for t in titles) and any("Solution" in t for t in titles)
    assert any(s.has_notes_slide and "say this" in s.notes_slide.notes_text_frame.text for s in p.slides)
    # 7 bullets must be split across slides, never overflowing one
    assert sum("Solution" in t for t in titles) >= 2


def test_csv_utf8_bom_and_numbers(doc, tmp_path):
    r = build_file("csv", doc, tmp_path, "t")
    raw = open(r.path, "rb").read()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(open(r.path, encoding="utf-8-sig")))
    assert rows[0] == ["Region", "Units", "Revenue"] and rows[1][0] == "North"


def test_xlsx_typed_cells(doc, tmp_path):
    r = build_file("xlsx", doc, tmp_path, "t")
    ws = load_workbook(r.path).worksheets[0]
    assert ws["A1"].value == "Region" and ws["B2"].value == 120 and isinstance(ws["C2"].value, (int, float))


def test_no_table_gives_clear_error(tmp_path):
    with pytest.raises(ValueError, match="no table"):
        build_file("csv", parse_markdown("# Just text\nhello"), tmp_path, "x")


def test_unique_names(doc, tmp_path):
    a = build_file("pdf", doc, tmp_path, "same")
    b = build_file("pdf", doc, tmp_path, "same")
    assert a.path != b.path


def test_unicode_and_indic_warning(tmp_path):
    r = build_file("pdf", parse_markdown("# வணக்கம்\nதமிழ் text — café ₹100"), tmp_path, "ta")
    assert r.warnings  # honest about complex-script shaping
    r2 = build_file("docx", parse_markdown("# வணக்கம்\nதமிழ்"), tmp_path, "ta")
    assert "வணக்கம்" in "".join(p.text for p in Document(r2.path).paragraphs)


def test_all_themes_readable():
    themes = catalog()
    assert len(themes) >= 140
    for t in themes.values():
        for slides in (False, True):
            p = document_palette(t, on_slides=slides)
            assert contrast(p.text, p.paper) >= 7, t.id
            assert contrast(p.heading, p.paper) >= 4.5, t.id


def test_unknown_theme_falls_back():
    assert get_theme("nope-nope").id == get_theme(None).id


def test_recommender_is_explained_and_respects_dark():
    r = recommend("dark developer dashboard", "pptx", limit=3)
    assert len(r) == 3 and all(x["reasons"] for x in r)
    assert any(x["colors"] for x in r)
    docs = recommend("formal legal contract", "pdf", limit=5)
    assert sum(1 for x in docs if get_theme(x["id"]).dark) <= 1
