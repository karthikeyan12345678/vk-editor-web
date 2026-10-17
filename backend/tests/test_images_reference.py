import io
import zipfile
from pathlib import Path

import pytest

from ltxps.assistant import images, reference
from ltxps.assistant.images import ImageConfig, ImageError


def _png() -> bytes:
    from PIL import Image

    b = io.BytesIO()
    Image.new("RGB", (64, 64), (10, 120, 200)).save(b, "PNG")
    return b.getvalue()


def test_default_engine_has_no_logo():
    assert images.image_config_from({}).provider == "horde"
    assert images.image_config_from({"provider": "pollinations"}).provider == "horde"  # old default is not used unless picked on purpose
    assert images.image_config_from({"provider": "pollinations", "chosen": True}).provider == "pollinations"


def test_busy_free_engine_falls_back_and_is_flagged(tmp_path, monkeypatch):
    monkeypatch.setattr(images, "_horde", lambda *a, **k: (_ for _ in ()).throw(ImageError("busy")))
    monkeypatch.setattr(images, "_pollinations", lambda *a, **k: _png())
    out = images.generate_image(ImageConfig("horde"), "a red car on a road", tmp_path)
    assert out["provider"] == "pollinations" and out["watermark"] is True


def test_horde_result_has_no_flag(tmp_path, monkeypatch):
    monkeypatch.setattr(images, "_horde", lambda *a, **k: _png())
    out = images.generate_image(ImageConfig("horde"), "a red car on a road", tmp_path)
    assert out["provider"] == "horde" and out["watermark"] is False


def test_both_fail_message(tmp_path, monkeypatch):
    monkeypatch.setattr(images, "_horde", lambda *a, **k: (_ for _ in ()).throw(ImageError("queue is full")))
    monkeypatch.setattr(images, "_pollinations", lambda *a, **k: (_ for _ in ()).throw(ImageError("quota")))
    with pytest.raises(ImageError) as e:
        images.generate_image(ImageConfig("horde"), "a red car on a road", tmp_path)
    assert "queue is full" in str(e.value) and "quota" in str(e.value)


def _docx(path: Path):
    from docx import Document
    from docx.shared import RGBColor

    d = Document()
    d.add_heading("Quarterly Report", 0)
    for h in ("Summary", "Results", "Risks", "Next steps"):
        p = d.add_heading(h, 1)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x12, 0x6B, 0xB0)
        d.add_paragraph("Revenue grew steadily during the quarter, driven by new customers in three regions and better retention across the board.")
    d.save(str(path))


def test_reference_profile_and_theme(tmp_path):
    p = tmp_path / "ref.docx"
    _docx(p)
    prof = reference.profile(p)
    assert prof["outline"][:3] == ["Quarterly Report", "Summary", "Results"]
    assert "126BB0" in [c.upper() for c in prof["colors"]]
    text = reference.instructions(prof, "pdf")
    assert "Summary → Results" in text.replace("Quarterly Report → ", "")
    th = reference.theme_from(prof)
    assert th is not None and th.colors["accent"].upper() == "#126BB0"
    from ltxps.assistant.documents import build_file, parse_markdown

    r = build_file("pdf", parse_markdown("# T\n\n## A\n\nhello world"), tmp_path, "out", th.id)
    assert Path(r.path).exists()


def test_reference_rejects_spreadsheets(tmp_path):
    p = tmp_path / "x.csv"
    p.write_text("a,b\n1,2\n")
    with pytest.raises(Exception):
        reference.profile(p)


def test_broken_local_model_falls_back_instead_of_crashing(tmp_path, monkeypatch):
    def boom(*a, **k):
        raise ImageError("The picture model on this PC failed (RuntimeError: no kernel image).")

    monkeypatch.setattr(images, "_local", boom)
    monkeypatch.setattr(images, "_horde", lambda *a, **k: _png())
    out = images.generate_image(ImageConfig("local", local_dir=str(tmp_path)), "a bus on a road", tmp_path)
    assert out["provider"] == "horde" and "no kernel image" in out["fallback_reason"]
